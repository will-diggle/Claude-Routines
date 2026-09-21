"""
bilinguist_excel_export.py
===========================
Pulls the current word_dictionary data from Supabase and writes it to local
Excel files (one per language) matching the original BB_Dictionary_*.xlsx format.

Run this locally whenever you want a fresh copy of the database as spreadsheets.

Usage:
    python bilinguist_excel_export.py [--lang LANG] [--output-dir PATH]

Options:
    --lang        Export only this language code (de/fr/es/it/sv/tr/ar).
    --output-dir  Directory to write .xlsx files to.
                  Defaults to ~/Downloads/Language Database spreadsheets/
"""

import argparse
import json
import os
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
except ImportError:
    print("ERROR: openpyxl not installed. Run: /usr/local/bin/python3.9 -m pip install openpyxl")
    sys.exit(1)

SCRIPT_DIR = Path(__file__).parent
ENV_FILE   = SCRIPT_DIR.parent / ".env"

DEFAULT_OUTPUT_DIR = Path.home() / "Downloads" / "Language Database spreadsheets"

LANGUAGES = {
    "de": "German",
    "fr": "French",
    "es": "Spanish",
    "it": "Italian",
    "sv": "Swedish",
    "tr": "Turkish",
    "ar": "Arabic",
}

WORD_TYPES = ["verb", "noun", "adjective", "adverb", "other"]

SHEET_NAMES = {
    "verb":      "Verbs",
    "noun":      "Nouns",
    "adjective": "Adjectives",
    "adverb":    "Adverbs",
    "other":     "Other",
}

# Common columns present in every sheet
COMMON_COLS = [
    "word_id", "language", "word", "lemma", "frequency_rank", "translation",
    "level", "ipa", "explanation", "example_sentence", "example_translation",
    "tip", "word_family", "common_collocations", "governed_prepositions",
    "source",
]


def _load_env() -> dict:
    env: dict = {}
    if not ENV_FILE.exists():
        return env
    for line in ENV_FILE.read_text().splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, _, v = line.partition("=")
        env[k.strip()] = v.strip()
    return env


# Reverse of bilinguist_excel_import.py's per-language tense-label -> column-prefix
# and pronoun-key -> column-suffix mappings. Kept in exact sync with that file's
# _tense_XX()/_build_data_XX_verb() functions on purpose: if these drift apart, a
# fresh export produces columns the import script won't recognize on the next
# round trip (the bug this replaces -- export used to derive column names from the
# raw DB tense label instead of matching import's fixed naming).
TENSE_LABEL_TO_PREFIX = {
    "de": {"PRÄSENS": "present", "PRÄTERITUM": "simplepast", "PERFEKT": "perfect",
           "PLUSQUAMPERFEKT": "pluperfect", "FUTUR I": "future1", "FUTUR II": "future2",
           "KONJUNKTIV I": "subj1", "KONJUNKTIV II": "subj2", "KONDITIONALIS": "conditional"},
    "fr": {"PRÉSENT": "present", "PASSÉ COMPOSÉ": "passe_compose", "IMPARFAIT": "imparfait",
           "PASSÉ SIMPLE": "passe_simple", "PLUS-QUE-PARFAIT": "plus_que_parfait", "FUTUR": "futur",
           "FUTUR ANTÉRIEUR": "futur_anterieur", "CONDITIONNEL": "conditionnel",
           "CONDITIONNEL PASSÉ": "conditionnel_passe", "SUBJONCTIF": "subjonctif",
           "SUBJONCTIF PASSÉ": "subjonctif_passe"},
    "es": {"PRESENTE": "presente", "PRETÉRITO INDEFINIDO": "indefinido", "IMPERFECTO": "imperfecto",
           "PRETÉRITO PERFECTO": "perfecto", "PLUSCUAMPERFECTO": "pluscuamperfecto", "FUTURO": "futuro",
           "FUTURO PERFECTO": "futuro_perfecto", "CONDICIONAL": "condicional",
           "CONDICIONAL COMPUESTO": "condicional_compuesto", "SUBJUNTIVO PRESENTE": "subjuntivo_presente",
           "SUBJUNTIVO IMPERFECTO": "subjuntivo_imperfecto", "SUBJUNTIVO FUTURO": "subjuntivo_futuro"},
    "it": {"PRESENTE": "presente", "PASSATO PROSSIMO": "passato_prossimo", "IMPERFETTO": "imperfetto",
           "PASSATO REMOTO": "passato_remoto", "TRAPASSATO PROSSIMO": "trapassato_prossimo", "FUTURO": "futuro",
           "FUTURO ANTERIORE": "futuro_anteriore", "CONDIZIONALE": "condizionale",
           "CONDIZIONALE PASSATO": "condizionale_passato", "CONGIUNTIVO PRESENTE": "congiuntivo_presente",
           "CONGIUNTIVO PASSATO": "congiuntivo_passato", "CONGIUNTIVO IMPERFETTO": "congiuntivo_imperfetto",
           "CONGIUNTIVO TRAPASSATO": "congiuntivo_trapassato"},
    "tr": {"GENİŞ ZAMAN": "geniszaman", "ŞİMDİKİ ZAMAN": "simdikizaman", "GELECEK ZAMAN": "gelecekzaman",
           "GEÇMİŞ ZAMAN (-DI)": "gecmiszaman", "ÖĞRENİLEN GEÇMİŞ (-MIŞ)": "misligecmis",
           "ŞART KİPİ": "sart", "GEREKLİLİK KİPİ": "gereklilik"},
    # Swedish verbs don't conjugate by person -- one plain column per tense (no
    # pronoun suffix), so this maps label directly to the full column name.
    "sv": {"PRESENS": "presens", "PRETERITUM": "preteritum", "PERFEKT": "perfekt",
           "PLUSKVAMPERFEKT": "pluskvamperfekt", "FUTURUM (ska)": "futurum_ska",
           "FUTURUM (kommer att)": "futurum_kommer_att", "KONDITIONALIS": "konditionalis"},
}

PRONOUN_TO_SUFFIX = {
    "de": {"ich": "ich", "du": "du", "er/sie/es": "er", "wir": "wir", "ihr": "ihr", "sie/Sie": "sie"},
    "fr": {"je": "je", "tu": "tu", "il/elle": "il", "nous": "nous", "vous": "vous", "ils/elles": "ils"},
    "es": {"yo": "yo", "tú": "tu", "él/ella": "el", "nosotros": "nosotros", "vosotros": "vosotros", "ellos/ellas": "ellos"},
    "it": {"io": "io", "tu": "tu", "lui/lei": "lui", "noi": "noi", "voi": "voi", "loro": "loro"},
    "tr": {"ben": "ben", "sen": "sen", "o": "o", "biz": "biz", "siz": "siz", "onlar": "onlar"},
}

IMPERATIVE_COL = {
    "de": {"du": "imperative_du", "ihr": "imperative_ihr", "Sie": "imperative_sie"},
    "fr": {"tu": "imperatif_tu", "nous": "imperatif_nous", "vous": "imperatif_vous"},
    "es": {
        "tú (afirm.)": "imperativo_afirmativo_tu", "usted (afirm.)": "imperativo_afirmativo_usted",
        "nosotros (afirm.)": "imperativo_afirmativo_nosotros", "vosotros (afirm.)": "imperativo_afirmativo_vosotros",
        "ustedes (afirm.)": "imperativo_afirmativo_ustedes", "tú (neg.)": "imperativo_negativo_tu",
        "usted (neg.)": "imperativo_negativo_usted", "nosotros (neg.)": "imperativo_negativo_nosotros",
        "vosotros (neg.)": "imperativo_negativo_vosotros", "ustedes (neg.)": "imperativo_negativo_ustedes",
    },
    "it": {"tu": "imperativo_tu", "noi": "imperativo_noi", "voi": "imperativo_voi"},
    "tr": {"sen": "imperative_sen", "siz": "imperative_siz"},
}

# Arabic: tense labels are Arabic-script dict keys, and each tense type uses a
# DIFFERENT pronoun subset (madi/mudari = full 13-person paradigm, subjunctive/
# jussive = the 8-person subset _build_data_ar_verb calls subj_pronouns).
AR_TENSE_LABEL_TO_PREFIX = {
    "الماضي": "madi", "المضارع المرفوع": "mudari",
    "المضارع المنصوب": "subjunctive", "المضارع المجزوم": "jussive",
}
AR_PRONOUN_SUFFIX_FULL = {
    "أنا": "ana", "أنتَ": "anta", "أنتِ": "anti", "هو": "huwa", "هي": "hiya",
    "أنتما": "antuma", "هما (م)": "huma_m", "هما (ف)": "huma_f", "نحن": "nahnu",
    "أنتم": "antum", "أنتن": "antunna", "هم": "hum", "هن": "hunna",
}
AR_PRONOUN_SUFFIX_SUBJ = {
    "أنا": "ana", "أنتَ": "anta", "أنتِ": "anti", "هو": "huwa", "هي": "hiya",
    "نحن": "nahnu", "أنتم": "antum", "هم": "hum",
}
AR_IMPERATIVE_COL = {
    "أنتَ": "imperative_anta", "أنتِ": "imperative_anti", "أنتما": "imperative_antuma",
    "أنتم": "imperative_antum", "أنتن": "imperative_antunna",
}


def _flatten_tenses(lang: str, tenses, flat: dict) -> None:
    """Write each tense's per-pronoun forms to their own column, using the exact
    column names bilinguist_excel_import.py expects for this language."""
    if not isinstance(tenses, dict):
        # Some rows (mixed generation pipelines -- see project memory on
        # word_dictionary sources) store a malformed shape here. Skip rather
        # than crash the whole export over one bad row's tenses data.
        return
    if lang == "sv":
        for label, table in tenses.items():
            col = TENSE_LABEL_TO_PREFIX.get("sv", {}).get(label)
            if col and isinstance(table, dict):
                flat[col] = table.get("—")
        return
    if lang == "ar":
        for label, table in tenses.items():
            prefix = AR_TENSE_LABEL_TO_PREFIX.get(label)
            if not prefix or not isinstance(table, dict):
                continue
            pronoun_map = AR_PRONOUN_SUFFIX_FULL if prefix in ("madi", "mudari") else AR_PRONOUN_SUFFIX_SUBJ
            for pronoun, form in table.items():
                suffix = pronoun_map.get(pronoun)
                if suffix:
                    flat[f"{prefix}_{suffix}"] = form
        return

    label_map = TENSE_LABEL_TO_PREFIX.get(lang, {})
    pronoun_map = PRONOUN_TO_SUFFIX.get(lang, {})
    for label, table in tenses.items():
        if not isinstance(table, dict):
            continue
        prefix = label_map.get(label)
        if not prefix:
            continue  # unrecognized tense label -- omit rather than guess a column name
        for pronoun, form in table.items():
            suffix = pronoun_map.get(pronoun)
            if suffix:
                flat[f"{prefix}_{suffix}"] = form


def _flatten_imperative(lang: str, imperative, flat: dict) -> None:
    """Same idea as _flatten_tenses, for the separately-stored imperative mood."""
    if lang == "sv":
        # Swedish stores imperative as a single plain string, not a per-pronoun dict.
        if isinstance(imperative, str) and imperative:
            flat["imperative"] = imperative
        return
    if not isinstance(imperative, dict):
        return  # malformed row -- see _flatten_tenses
    if lang == "ar":
        for pron, form in imperative.items():
            col = AR_IMPERATIVE_COL.get(pron)
            if col:
                flat[col] = form
        return
    col_map = IMPERATIVE_COL.get(lang, {})
    for pron, form in imperative.items():
        col = col_map.get(pron)
        if col:
            flat[col] = form


def _flatten_row(row: dict) -> dict:
    """Flatten a DB row (common cols + JSONB data) into a flat dict for Excel."""
    flat = {
        "word_id":              row.get("word_id"),
        "language":             row.get("language"),
        "word":                 row.get("word"),
        "lemma":                row.get("lemma"),
        "frequency_rank":       row.get("frequency_rank"),
        "translation":          row.get("translation"),
        "level":                row.get("level"),
        "ipa":                  row.get("ipa"),
        "explanation":          row.get("explanation"),
        "example_sentence":     row.get("example_sentence"),
        "example_translation":  row.get("example_translation"),
        "tip":                  row.get("tip"),
        "word_family":          row.get("word_family"),
        "common_collocations":  row.get("common_collocations"),
        "governed_prepositions": row.get("governed_prepositions"),
        "source":               row.get("source"),
    }

    data = row.get("data") or {}
    lang = row.get("language")

    # Flatten all data keys except the ones handled specially below. Some rows
    # (mixed generation pipelines -- different pipelines use different schemas,
    # e.g. a camelCase nested object instead of flat snake_case fields) carry a
    # dict/list where a scalar is expected. Excel cells can't hold those directly,
    # so serialize to JSON text rather than dropping the data or crashing the export.
    for k, v in data.items():
        if k in ("tenses", "cases", "imperative", "declension", "possessive"):
            continue
        if isinstance(v, (dict, list)):
            v = json.dumps(v, ensure_ascii=False)
        flat[k] = v

    # Tenses and imperative use language-specific column naming that must match
    # bilinguist_excel_import.py exactly -- see _flatten_tenses/_flatten_imperative.
    _flatten_tenses(lang, data.get("tenses") or {}, flat)
    _flatten_imperative(lang, data.get("imperative"), flat)

    # Flatten cases (column names already match import's raw column keys 1:1).
    # Guarded against malformed rows the same way as tenses/imperative above.
    cases = data.get("cases")
    if isinstance(cases, dict):
        for case_col, val in cases.items():
            flat[case_col] = val

    # Flatten declension (German adjectives)
    declension = data.get("declension")
    if isinstance(declension, dict):
        for dcol, val in declension.items():
            flat[dcol] = val

    # Flatten possessive (Turkish nouns)
    possessive = data.get("possessive")
    if isinstance(possessive, dict):
        for p_key, val in possessive.items():
            flat[f"possessive_{p_key}"] = val

    return flat


def _write_sheet(ws, rows: list[dict]):
    """Write rows to a worksheet, auto-detecting columns from all rows."""
    if not rows:
        ws.append(["(no data)"])
        return

    # Collect all keys, common cols first
    all_keys: list[str] = list(COMMON_COLS)
    extra_keys: set[str] = set()
    for row in rows:
        for k in row:
            if k not in all_keys and k not in extra_keys:
                extra_keys.add(k)
    all_keys.extend(sorted(extra_keys))

    # Header row
    ws.append(all_keys)
    header_row = ws[1]
    header_fill = PatternFill("solid", fgColor="CCDDEE")
    header_font = Font(bold=True)
    for cell in header_row:
        cell.fill = header_fill
        cell.font = header_font

    # Data rows
    for row in rows:
        ws.append([row.get(k) for k in all_keys])

    # Freeze header
    ws.freeze_panes = "A2"


def export_language(supa, lang: str, output_dir: Path):
    lang_name = LANGUAGES.get(lang, lang)
    output_path = output_dir / f"BB_Dictionary_{lang_name}.xlsx"

    print(f"[export] Exporting {lang_name}...")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default empty sheet

    for word_type in WORD_TYPES:
        sheet_name = SHEET_NAMES[word_type]

        # Fetch from Supabase
        try:
            resp = (
                supa.table("word_dictionary")
                .select("*")
                .eq("language", lang)
                .eq("word_type", word_type)
                .order("frequency_rank", desc=False)
                .execute()
            )
            db_rows = resp.data or []
        except Exception as e:
            print(f"  [WARN] Failed to fetch {lang}/{word_type}: {e}")
            db_rows = []

        flat_rows = [_flatten_row(r) for r in db_rows]
        ws = wb.create_sheet(title=sheet_name)
        _write_sheet(ws, flat_rows)
        print(f"  [{sheet_name}] {len(flat_rows)} rows")

    wb.save(output_path)
    print(f"  Saved: {output_path}")


def main():
    parser = argparse.ArgumentParser(description="Export Supabase word_dictionary to Excel")
    parser.add_argument("--lang",       help="Only export this language code")
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR), help="Output directory")
    args = parser.parse_args()

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    env = _load_env()
    supa_url = env.get("EXPO_PUBLIC_SUPABASE_URL") or os.getenv("EXPO_PUBLIC_SUPABASE_URL", "")
    supa_key = env.get("SUPABASE_SERVICE_ROLE_KEY") or os.getenv("SUPABASE_SERVICE_ROLE_KEY", "")

    if not supa_url or not supa_key:
        print("[ERROR] Missing EXPO_PUBLIC_SUPABASE_URL or SUPABASE_SERVICE_ROLE_KEY", file=sys.stderr)
        sys.exit(1)

    from supabase import create_client
    supa = create_client(supa_url, supa_key)
    print(f"[export] Supabase connected. Output: {output_dir}")

    langs = [args.lang] if args.lang else list(LANGUAGES.keys())
    for lang in langs:
        export_language(supa, lang, output_dir)

    print("[export] Done.")


if __name__ == "__main__":
    main()
